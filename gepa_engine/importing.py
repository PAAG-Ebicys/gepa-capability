"""Read the cases of a preparation request and record where each one came from.

Inputs are cases written inline (for example, proposed in a conversation), cells pasted from a
spreadsheet (tab-separated text), or a JSON, CSV, TSV or ``.xlsx`` file. The engine records what
it observes (file name and SHA-256, sheet, row); what it cannot observe (where pasted cells or
inline cases come from) the caller declares in ``origin``. Nothing here invents, completes or
relabels a case: table columns only map to case fields by their header.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from collections.abc import Callable, Iterable
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from .datasets import MAX_CASES, fold

MAX_INPUTS = 50
MAX_ORIGIN = 500
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_BYTES = 40 * 1024 * 1024  # uncompressed size of an .xlsx
MAX_ARCHIVE_ENTRIES = 2000
MAX_SHEETS = 20
MAX_COLUMNS = 100
FORMATS = {".json": "json", ".csv": "csv", ".tsv": "tsv", ".txt": "tsv", ".xlsx": "xlsx"}
DELIMITERS = (",", ";", "\t")  # Excel writes ';' in locales where ',' is the decimal separator
# Header (compared without case, accents or extra spaces) -> case field.
HEADERS = {
    "id": "id", "identificador": "id",
    "entrada": "input", "input": "input", "estado": "input", "state": "input",
    "esperado": "expected", "respuesta esperada": "expected", "opcion esperada": "expected", "expected": "expected",
    "conjunto": "split", "particion": "split", "split": "split",
    "procedencia": "source", "origen": "source", "fuente": "source", "source": "source",
    "grupo": "group", "group": "group",
    "contexto": "context", "context": "context",
    "criterio": "criterion", "criterion": "criterion",
    "opciones": "options", "options": "options",
    "escenario": "scenario", "scenario": "scenario",
    "comprobacion": "check", "check": "check",
    "entorno": "environment", "environment": "environment",
    "notas": "notes", "nota": "notes", "notes": "notes",
}

Locate = Callable[[int], dict[str, Any]]
Rows = list[tuple[int, list[str]]]  # (row number as the spreadsheet shows it, cells)


class InputError(ValueError):
    """User-facing reason an input of the request cannot be read."""


def load_inputs(inputs: Any, base_dir: Path | str | None) -> list[dict[str, Any]]:
    """Every case of every input, in order, each with a ``source`` provenance object."""
    if not isinstance(inputs, list) or not 1 <= len(inputs) <= MAX_INPUTS:
        raise InputError(f"'inputs' debe ser una lista de 1 a {MAX_INPUTS} entradas.")
    cases: list[dict[str, Any]] = []
    for number, entry in enumerate(inputs, 1):
        cases.extend(_load(entry, number, base_dir))
        if len(cases) > MAX_CASES:
            raise InputError(f"El dataset admite como máximo {MAX_CASES} casos.")
    return cases


def _load(entry: Any, number: int, base_dir: Path | str | None) -> list[dict[str, Any]]:
    label = f"entrada {number}"
    if not isinstance(entry, dict) or sum(key in entry for key in ("path", "table", "cases")) != 1:
        raise InputError(f"La {label} debe tener exactamente uno de 'path', 'table' o 'cases'.")
    if "cases" in entry:
        _allowed(entry, {"cases", "origin"}, label)
        return _from_documents(entry["cases"], lambda index: {"kind": "inline", "index": index}, _origin(entry, label), label, inline=True)
    if "table" in entry:
        _allowed(entry, {"table", "origin"}, label)
        origin = _origin(entry, label)
        if origin is None:
            raise InputError(f"Indica 'origin' en la {label}: de dónde vienen las celdas pegadas (por ejemplo, «hoja de seguimiento del equipo, copiada por la persona»).")
        table = entry["table"]
        if not isinstance(table, str) or len(table.encode("utf-8")) > MAX_FILE_BYTES:
            raise InputError(f"'table' de la {label} debe ser texto de hasta 10 MB.")
        return _from_rows(_text_rows(table, "\t", label), lambda row: {"kind": "pasted", "format": "tsv", "row": row}, origin, label)
    _allowed(entry, {"path", "format", "sheet", "origin"}, label)
    path, raw = _read(entry["path"], base_dir, label)
    kind = entry.get("format") or FORMATS.get(path.suffix.lower())
    if kind not in set(FORMATS.values()):
        raise InputError(f"No se reconoce el formato de {path.name}: usa .xlsx, .csv, .tsv o .json (un libro .xls antiguo debe guardarse como .xlsx).")
    if "sheet" in entry and (kind != "xlsx" or not isinstance(entry["sheet"], str)):
        raise InputError(f"'sheet' de la {label} solo se indica, como texto, para un libro .xlsx.")
    base = {"kind": "file", "format": kind, "file": path.name, "fileSha256": hashlib.sha256(raw).hexdigest()}
    origin = _origin(entry, label)
    if kind == "json":
        return _from_documents(_json_cases(raw, path), lambda index: {**base, "index": index}, origin, label, inline=False)
    if kind == "xlsx":
        sheet, rows = _workbook_rows(raw, entry.get("sheet"), path)
        return _from_rows(rows, lambda row: {**base, "sheet": sheet, "row": row}, origin, label)
    text = _decoded(raw, path)
    return _from_rows(_text_rows(text, "\t" if kind == "tsv" else _delimiter(text), label), lambda row: {**base, "row": row}, origin, label)


def _allowed(entry: dict[str, Any], keys: set[str], label: str) -> None:
    unknown = set(entry) - keys
    if unknown:
        raise InputError(f"Campos no admitidos en la {label}: {', '.join(sorted(unknown))}.")


def _origin(entry: dict[str, Any], label: str) -> str | None:
    origin = entry.get("origin")
    if origin is None:
        return None
    if not isinstance(origin, str) or not origin.strip() or len(origin) > MAX_ORIGIN:
        raise InputError(f"'origin' de la {label} debe ser texto de hasta {MAX_ORIGIN} caracteres.")
    return origin.strip()


def _provenance(located: dict[str, Any], origin: str | None, declared: Any, where: str) -> dict[str, Any]:
    """Where the engine found the case, plus what the person declared about it (never mixed up)."""
    if declared is not None and (not isinstance(declared, str) or not declared.strip() or len(declared) > MAX_ORIGIN):
        raise InputError(f"La procedencia de {where} debe ser texto de hasta {MAX_ORIGIN} caracteres.")
    source = dict(located)
    if origin:
        source["origin"] = origin
    if declared:
        source["declared"] = declared.strip()
    return source


def _from_documents(items: Any, locate: Locate, origin: str | None, label: str, *, inline: bool) -> list[dict[str, Any]]:
    """Cases given as JSON objects; a case's own ``source`` text is kept as its declared provenance."""
    if not isinstance(items, list):
        raise InputError(f"La {label} debe contener una lista de casos.")
    cases = []
    for index, item in enumerate(items, 1):
        if not isinstance(item, dict):
            raise InputError(f"El caso {index} de la {label} debe ser un objeto JSON.")
        case = dict(item)
        declared = case.pop("source", None)
        if inline and declared is None and origin is None:
            raise InputError(f"Indica 'origin' en la {label}: de dónde vienen sus casos (por ejemplo, «propuestos en la conversación y confirmados por la persona»).")
        case["source"] = _provenance(locate(index), origin, declared, f"el caso {index} de la {label}")
        cases.append(case)
    return cases


def _from_rows(rows: Rows, locate: Locate, origin: str | None, label: str) -> list[dict[str, Any]]:
    """Cases from a table whose first non-empty row names the columns."""
    rows = [(number, cells) for number, cells in rows if any(cell.strip() for cell in cells)]
    if not rows:
        raise InputError(f"La {label} no contiene filas con datos.")
    _, header = rows[0]
    while header and not header[-1].strip():
        header = header[:-1]
    keys = [HEADERS.get(fold(cell)) for cell in header]
    unknown = [cell for cell, key in zip(header, keys) if key is None]
    if unknown:
        raise InputError(f"Encabezados no reconocidos en la {label}: {', '.join(repr(cell) for cell in unknown)}. "
                         "Usa entrada, esperado, conjunto, id, procedencia, grupo, contexto, criterio, opciones o notas.")
    fields = [key for key in keys if key is not None]
    repeated = sorted({key for key in fields if fields.count(key) > 1})
    if repeated:
        raise InputError(f"La {label} repite columnas para: {', '.join(repeated)}.")
    if "input" not in fields:
        raise InputError(f"La {label} necesita una columna 'entrada'.")
    cases = []
    for number, cells in rows[1:]:
        if any(cell.strip() for cell in cells[len(fields):]):
            raise InputError(f"La fila {number} de la {label} tiene datos fuera de las columnas con encabezado.")
        # An empty cell is an absent field, except in an 'id' column, where it is reported.
        case: dict[str, Any] = {key: cell.strip() for key, cell in zip(fields, cells) if cell.strip() or key == "id"}
        case["source"] = _provenance(locate(number), origin, case.pop("source", None), f"la fila {number} de la {label}")
        cases.append(case)
    return cases


def _read(value: Any, base_dir: Path | str | None, label: str) -> tuple[Path, bytes]:
    if not isinstance(value, str) or not value.strip():
        raise InputError(f"'path' de la {label} debe ser una ruta.")
    path = Path(value).expanduser()
    if not path.is_absolute():
        if base_dir is None:
            raise InputError(f"'path' de la {label} debe ser absoluta cuando la solicitud no viene de un archivo.")
        path = Path(base_dir) / path
    path = path.resolve()
    try:
        if path.stat().st_size > MAX_FILE_BYTES:
            raise InputError(f"{path.name} supera el límite de 10 MB.")
        return path, path.read_bytes()
    except OSError:
        raise InputError(f"No se pudo leer {path}.") from None


def _json_cases(raw: bytes, path: Path) -> Any:
    try:
        document = json.loads(raw.decode("utf-8-sig"))
    except ValueError:
        raise InputError(f"{path.name} no contiene JSON válido en UTF-8.") from None
    if isinstance(document, dict):
        if "cases" not in document:
            raise InputError(f"{path.name} debe ser una lista de casos o un objeto con 'cases'.")
        return document["cases"]
    return document


def _decoded(raw: bytes, path: Path) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise InputError(f"{path.name} no está en UTF-8: en Excel, guárdalo como «CSV UTF-8 (delimitado por comas)».") from None


def _delimiter(text: str) -> str:
    """The separator that makes the most header cells recognizable (quoted or not); ',' when none does."""
    first = next((line for line in text.splitlines() if line.strip()), "")

    def recognized(delimiter: str) -> int:
        return sum(fold(cell) in HEADERS for cell in next(csv.reader([first], delimiter=delimiter), []))

    return max(DELIMITERS, key=recognized)


def _text_rows(text: str, delimiter: str, label: str) -> Rows:
    csv.field_size_limit(MAX_FILE_BYTES)
    try:
        rows = list(enumerate(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True), 1))
    except csv.Error:
        raise InputError(f"La {label} tiene comillas o columnas mal formadas.") from None
    return _bounded(rows, label)


def _bounded(rows: Iterable[tuple[int, list[str]]], label: str) -> Rows:
    bounded: Rows = []
    for number, cells in rows:
        if len(bounded) > MAX_CASES or len(cells) > MAX_COLUMNS:
            raise InputError(f"La {label} supera {MAX_CASES} filas o {MAX_COLUMNS} columnas.")
        bounded.append((number, cells))
    return bounded


def _workbook_rows(raw: bytes, requested: str | None, path: Path) -> tuple[str, Rows]:
    """Cell values of one sheet as text. Formulas, external links and XML entities are refused."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        raise InputError(f"{path.name} no es un libro .xlsx válido.") from None
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES or sum(item.file_size for item in entries) > MAX_ARCHIVE_BYTES:
            raise InputError(f"{path.name} supera 40 MB descomprimido o contiene demasiados archivos.")
        if "xl/workbook.xml" not in archive.namelist():
            raise InputError(f"{path.name} no es un libro .xlsx válido.")
        for item in entries:
            if item.filename.endswith((".xml", ".rels")):
                content = archive.read(item).upper()
                if b"<!DOCTYPE" in content or b"<!ENTITY" in content:
                    raise InputError(f"{path.name} contiene XML no compatible; guárdalo de nuevo como .xlsx.")
    from openpyxl import load_workbook  # type: ignore[import-untyped]  # imported here: only an Excel import needs it

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Exception:
        raise InputError(f"No se pudo leer {path.name}; usa un .xlsx sin contraseña.") from None
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise InputError(f"{path.name} tiene más de {MAX_SHEETS} hojas.")
        sheet = _choose_sheet(workbook.worksheets, requested, path)
        return sheet.title, _bounded(((number, _cells(row, sheet.title)) for number, row in enumerate(sheet.iter_rows(min_row=1), 1)), f"hoja «{sheet.title}»")
    finally:
        workbook.close()


def _choose_sheet(sheets: list[Any], requested: str | None, path: Path) -> Any:
    names = ", ".join(sheet.title for sheet in sheets)
    if requested is not None:
        match = next((sheet for sheet in sheets if fold(sheet.title) == fold(requested)), None)
        if match is None:
            raise InputError(f"{path.name} no tiene la hoja «{requested}»; hojas: {names}.")
        return match
    cases = [sheet for sheet in sheets if fold(sheet.title) == "casos"]  # the sheet name of the downloadable templates
    if cases:
        return cases[0]
    if len(sheets) == 1:
        return sheets[0]
    raise InputError(f"{path.name} tiene varias hojas ({names}); indica en 'sheet' cuál contiene los casos.")


def _cells(row: Iterable[Any], title: str) -> list[str]:
    values = []
    for cell in row:
        if cell.data_type == "f":
            raise InputError(f"La hoja «{title}» tiene una fórmula en {cell.coordinate}: pega sus resultados como valores antes de importar.")
        value = cell.value
        values.append("" if value is None else value.isoformat() if isinstance(value, (datetime, date, time)) else str(value))
    return values
